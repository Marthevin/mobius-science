#!/usr/bin/env python3
"""Inspect a scientific DOCX and optionally render its exact bytes for page QA."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import sys
import tempfile
from dataclasses import asdict, dataclass
from pathlib import Path
from urllib.parse import unquote
from zipfile import BadZipFile, ZipFile
from xml.sax.saxutils import escape

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn


BUNDLED_CJK_FONT = (
    Path(__file__).resolve().parents[1] / "assets/fonts/NotoSerifSC-Regular.otf"
)


@dataclass(frozen=True)
class Finding:
    severity: str
    code: str
    message: str


def _doi_set(value: str) -> set[str]:
    found = set()
    for match in re.finditer(r"\b10\.\d{4,9}/[^\s<>]+", unquote(value), re.I):
        doi = match.group(0).rstrip(".,;:，。；：")
        while doi.endswith(")") and doi.count(")") > doi.count("("):
            doi = doi[:-1]
        found.add(doi.casefold())
    return found


def file_identity(path: Path) -> dict[str, int | str]:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return {"bytes": path.stat().st_size, "sha256": digest.hexdigest()}


def inspect_docx(
    path: Path,
    *,
    language: str = "mixed",
    minimum_words: int = 0,
    minimum_cjk_chars: int = 0,
    require_doi_links: bool = False,
) -> tuple[dict[str, int], list[Finding]]:
    """Read the actual OOXML package, not a converted preview."""

    findings: list[Finding] = []
    try:
        with ZipFile(path) as package:
            corrupt_part = package.testzip()
            if corrupt_part:
                return {}, [Finding("error", "corrupt-package", corrupt_part)]
        document = Document(path)
    except (BadZipFile, KeyError, ValueError, OSError) as exc:
        return {}, [Finding("error", "invalid-docx", str(exc))]

    paragraphs = list(document.paragraphs)
    text_parts = [paragraph.text for paragraph in paragraphs]
    for table in document.tables:
        text_parts.extend(cell.text for row in table.rows for cell in row.cells)
    text = "\n".join(text_parts)
    english_words = len(re.findall(r"\b[A-Za-z][A-Za-z'’-]*\b", text))
    cjk_characters = len(re.findall(r"[\u3400-\u9fff]", text))
    metrics = {
        "english_words": english_words,
        "cjk_characters": cjk_characters,
        "tables": len(document.tables),
        "figures": len(document.inline_shapes),
        "headings": sum(p.style.name.startswith("Heading ") for p in paragraphs),
    }

    if not any(p.style.name == "Title" and p.text.strip() for p in paragraphs):
        findings.append(
            Finding("warning", "missing-title", "No non-empty Word Title paragraph")
        )
    if language == "en" and cjk_characters >= 20:
        findings.append(
            Finding(
                "warning",
                "unexpected-cjk",
                "English report contains substantial CJK prose",
            )
        )
    if language == "zh" and minimum_words:
        minimum_cjk_chars = max(minimum_cjk_chars, minimum_words)
    elif minimum_words and english_words < minimum_words:
        findings.append(
            Finding(
                "warning",
                "short-report",
                f"English word count {english_words} is below {minimum_words}",
            )
        )
    if minimum_cjk_chars and cjk_characters < minimum_cjk_chars:
        findings.append(
            Finding(
                "warning",
                "short-report-cjk",
                f"CJK character count {cjk_characters} is below {minimum_cjk_chars}",
            )
        )

    table_captions = sum(
        p.style.name == "Report Caption"
        and re.match(r"^(?:Table|表)\s*\d+", p.text, re.I) is not None
        for p in paragraphs
    )
    figure_captions = sum(
        p.style.name == "Report Caption"
        and re.match(r"^(?:Figure|图)\s*\d+", p.text, re.I) is not None
        for p in paragraphs
    )
    if table_captions < len(document.tables):
        findings.append(
            Finding("warning", "uncaptioned-table", "A table lacks a numbered caption")
        )
    if figure_captions < len(document.inline_shapes):
        findings.append(
            Finding(
                "warning", "uncaptioned-figure", "A figure lacks a numbered caption"
            )
        )

    if require_doi_links:
        visible = _doi_set(text)
        linked = set()
        for relationship in document.part.rels.values():
            if relationship.reltype == RT.HYPERLINK and relationship.is_external:
                target = relationship.target_ref
                if re.match(r"https?://(?:dx\.)?doi\.org/", target, re.I):
                    linked.update(_doi_set(target))
        missing = sorted(visible - linked)
        if missing:
            findings.append(
                Finding(
                    "warning",
                    "unlinked-doi",
                    f"{len(missing)} visible DOI(s) have no matching hyperlink: {', '.join(missing[:3])}",
                )
            )
    return metrics, findings


def fontconfig_environment(output_dir: Path) -> dict[str, str]:
    """Expose the bundled CJK font to LibreOffice without installing it system-wide."""

    environment = os.environ.copy()
    if sys.platform == "win32" or not BUNDLED_CJK_FONT.is_file():
        return environment
    directories = [BUNDLED_CJK_FONT.parent]
    if sys.platform == "darwin":
        directories.extend(
            Path(path)
            for path in (
                "/System/Library/Fonts",
                "/System/Library/Fonts/Supplemental",
                "/Library/Fonts",
            )
        )
        directories.append(Path.home() / "Library/Fonts")
    else:
        directories.extend(
            Path(path) for path in ("/usr/share/fonts", "/usr/local/share/fonts")
        )
        directories.append(Path.home() / ".local/share/fonts")
    cache = output_dir / "fontcache"
    cache.mkdir(parents=True, exist_ok=True)
    config = output_dir / "fontconfig.xml"
    lines = [
        '<?xml version="1.0"?>',
        '<!DOCTYPE fontconfig SYSTEM "fonts.dtd">',
        "<fontconfig>",
        *(
            f"  <dir>{escape(str(directory))}</dir>"
            for directory in directories
            if directory.is_dir()
        ),
        f"  <cachedir>{escape(str(cache))}</cachedir>",
        "</fontconfig>",
    ]
    config.write_text("\n".join(lines) + "\n", encoding="utf-8")
    environment["FONTCONFIG_FILE"] = str(config)
    return environment


def missing_expected_cjk_font(pdf_fonts: str, expected_font: str) -> bool:
    if expected_font != "Noto Serif SC":
        return False
    return "NotoSerifSC-" not in pdf_fonts


def render_docx(path: Path, output_dir: Path, soffice: Path) -> dict[str, object]:
    """Convert the exact DOCX with an isolated LibreOffice profile, then rasterize."""

    output_dir.mkdir(parents=True, exist_ok=True)
    proof = output_dir / f"{path.stem}.pdf"
    proof.unlink(missing_ok=True)
    with tempfile.TemporaryDirectory(prefix="docx-qa-lo-") as profile:
        command = [
            str(soffice),
            f"-env:UserInstallation={Path(profile).as_uri()}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(path),
        ]
        conversion = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=120,
            env=fontconfig_environment(output_dir),
        )
    if conversion.returncode != 0 or not proof.is_file() or proof.stat().st_size == 0:
        details = (
            conversion.stderr
            or conversion.stdout
            or f"exit code {conversion.returncode}; no PDF produced"
        )
        raise RuntimeError(f"DOCX render failed: {details}")
    pages_dir = output_dir / "pages"
    pages_dir.mkdir(exist_ok=True)
    for prior in pages_dir.glob("page-*.png"):
        prior.unlink()
    raster = subprocess.run(
        ["pdftoppm", "-png", "-r", "120", str(proof), str(pages_dir / "page")],
        capture_output=True,
        text=True,
        timeout=120,
    )
    if raster.returncode != 0:
        raise RuntimeError(f"Page rasterization failed: {raster.stderr}")
    pages = sorted(
        pages_dir.glob("page-*.png"), key=lambda page: int(page.stem.rsplit("-", 1)[-1])
    )
    if not pages:
        raise RuntimeError("DOCX render produced no pages")
    return {
        "proof_pdf": str(proof),
        "proof_identity": file_identity(proof),
        "page_count": len(pages),
        "pages": [str(page) for page in pages],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--language", choices=("en", "zh", "mixed"), default="mixed")
    parser.add_argument("--min-words", type=int, default=0)
    parser.add_argument("--min-cjk-chars", type=int, default=0)
    parser.add_argument("--require-doi-links", action="store_true")
    parser.add_argument(
        "--soffice", type=Path, help="LibreOffice binary for page rendering"
    )
    parser.add_argument("--strict", action="store_true", help="Fail on warnings")
    args = parser.parse_args()
    if args.min_words < 0 or args.min_cjk_chars < 0:
        parser.error("length floors must be nonnegative")
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    report: dict[str, object] = {"docx": str(args.docx.resolve())}
    if not args.docx.is_file():
        findings = [Finding("error", "missing-docx", "DOCX file not found")]
    else:
        report["file_identity"] = file_identity(args.docx)
        metrics, findings = inspect_docx(
            args.docx,
            language=args.language,
            minimum_words=args.min_words,
            minimum_cjk_chars=args.min_cjk_chars,
            require_doi_links=args.require_doi_links,
        )
        report["text_metrics"] = metrics
        if args.soffice:
            try:
                report["render"] = render_docx(args.docx, output_dir, args.soffice)
                if metrics.get("cjk_characters", 0):
                    document = Document(args.docx)
                    fonts = document.styles["Normal"].element.get_or_add_rPr().rFonts
                    expected = (
                        fonts.get(qn("w:eastAsia")) if fonts is not None else None
                    )
                    if expected == "Noto Serif SC":
                        font_result = subprocess.run(
                            ["pdffonts", str(report["render"]["proof_pdf"])],
                            capture_output=True,
                            text=True,
                            timeout=30,
                        )
                        if font_result.returncode or missing_expected_cjk_font(
                            font_result.stdout, expected
                        ):
                            findings.append(
                                Finding(
                                    "error",
                                    "cjk-font-missing-in-proof",
                                    "Rendered PDF did not embed Noto Serif SC; inspect Chinese glyphs",
                                )
                            )
            except (OSError, RuntimeError, subprocess.TimeoutExpired) as exc:
                findings.append(Finding("error", "render-failed", str(exc)))
        else:
            findings.append(
                Finding("warning", "visual-qa-pending", "Render and inspect every page")
            )
    report["findings"] = [asdict(item) for item in findings]
    (output_dir / "quality-report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    for item in findings:
        print(f"{item.severity.upper()} {item.code}: {item.message}")
    print(f"QA report: {output_dir / 'quality-report.json'}")
    return int(any(item.severity == "error" or args.strict for item in findings))


if __name__ == "__main__":
    raise SystemExit(main())
