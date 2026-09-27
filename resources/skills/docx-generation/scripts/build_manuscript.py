"""Deterministic prose-first Markdown export through the document Skills' templates.

Requires mistune>=3,<4 and the chosen format's dependencies. No network access,
source verification, rewriting, or automatic claim of visual QA is performed.
Unsupported syntax fails explicitly instead of silently dropping scholarly content.
"""
from __future__ import annotations

import argparse
import hashlib
import html
import importlib.util
import json
from pathlib import Path
import re
from urllib.parse import urlsplit
from xml.sax.saxutils import escape

import mistune


def spans(tokens, bold=False, italic=False, url=None):
    result = []
    for token in tokens:
        kind = token['type']
        if kind == 'text':
            text = html.unescape(token['raw'])
            start = 0
            for match in ([] if url else re.finditer(r'\b10\.\d{4,9}/[^\s<>]+', text)):
                identifier = match[0].rstrip('.,;:，。；：')
                while identifier.endswith(')') and identifier.count(')') > identifier.count('('):
                    identifier = identifier[:-1]
                result.append({'text': text[start:match.start()], 'bold': bold, 'italic': italic, 'url': None})
                result.append({'text': identifier, 'bold': bold, 'italic': italic, 'url': 'https://doi.org/' + identifier})
                start = match.start() + len(identifier)
            result.append({'text': text[start:], 'bold': bold, 'italic': italic, 'url': url})
        elif kind in {'softbreak', 'linebreak'}:
            result.append({'text': ' ', 'bold': bold, 'italic': italic, 'url': url})
        elif kind in {'strong', 'emphasis', 'link'}:
            target = token.get('attrs', {}).get('url', url)
            if target and (urlsplit(target).scheme not in {'http', 'https'} or not urlsplit(target).hostname):
                raise ValueError('Only explicit HTTP(S) document hyperlinks are supported')
            result.extend(spans(token['children'], bold or kind == 'strong', italic or kind == 'emphasis', target))
        else:
            raise ValueError(f'Unsupported inline syntax: {kind}; use the document template for complex content')
    return result


def plain(parts):
    return ''.join(part['text'] for part in parts)


def parse(text):
    tokens = [t for t in mistune.create_markdown(renderer='ast', plugins=['table', 'url'])(text) if t['type'] != 'blank_line']
    if not tokens or tokens[0]['type'] != 'heading' or tokens[0]['attrs']['level'] != 1:
        raise ValueError('Begin the manuscript with one # Title')
    title = spans(tokens.pop(0)['children'])
    blocks = []
    references = False
    for token in tokens:
        kind = token['type']
        if kind == 'heading':
            level = token['attrs']['level'] - 1
            if not 1 <= level <= 3:
                raise ValueError('Use one # title and ## through #### headings')
            content = spans(token['children'])
            references = plain(content).strip().casefold() in {'references', 'bibliography', '参考文献', '參考文獻'}
            blocks.append({'kind': 'heading', 'level': level, 'parts': content})
        elif kind == 'paragraph':
            if references and any(t['type'] in {'softbreak', 'linebreak'} for t in token['children']):
                raise ValueError('Put a blank line between references and keep each entry on one source line')
            blocks.append({'kind': 'reference' if references else 'paragraph', 'parts': spans(token['children'])})
        elif kind == 'block_quote':
            for child in token['children']:
                if child['type'] != 'paragraph':
                    raise ValueError('Use plain paragraphs inside a block quote')
                blocks.append({'kind': 'quote', 'parts': spans(child['children'])})
        elif kind == 'list':
            start = token.get('attrs', {}).get('start', 1)
            for index, item in enumerate(token['children'], start):
                if len(item['children']) != 1 or item['children'][0]['type'] not in {'block_text', 'paragraph'}:
                    raise ValueError('Nested or multi-paragraph lists require the document template')
                prefix = f'{index}. ' if token['attrs']['ordered'] else '• '
                blocks.append({'kind': 'reference' if references else 'paragraph', 'parts': [dict(text=prefix, bold=False, italic=False, url=None)] + spans(item['children'][0]['children'])})
        elif kind == 'table':
            if not blocks or blocks[-1]['kind'] != 'paragraph' or not re.match(r'^(Table|表)\s*\d+', plain(blocks[-1]['parts']), re.I):
                raise ValueError('A table needs a preceding numbered caption paragraph')
            caption = blocks.pop()['parts']
            head, body = token['children']
            rows = [[spans(cell['children']) for cell in head['children']]]
            rows += [[spans(cell['children']) for cell in row['children']] for row in body['children']]
            blocks.append({'kind': 'table', 'rows': rows, 'caption': caption})
        else:
            raise ValueError(f'Unsupported block syntax: {kind}; preserve it using the document template')
    return title, blocks


def load_template(path):
    spec = importlib.util.spec_from_file_location('document_template', path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def word_runs(paragraph, parts):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    for part in parts:
        run = paragraph.add_run(part['text'])
        run.bold, run.italic = True if part['bold'] else None, True if part['italic'] else None
        if part['url']:
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(part['url'], RT.HYPERLINK, is_external=True))
            link.append(run._r)
            paragraph._p.append(link)


def build_docx(text, output, template, language='en', latin_font='Times New Roman', cjk_font='Noto Serif SC'):
    title, blocks = parse(text)
    module = load_template(template)
    report = module.ScientificDocxReport(output, plain(title), language=language, profile='academic', latin_font=latin_font, cjk_font=cjk_font)
    report.add_front_matter()
    word_runs(report.document.paragraphs[0].clear(), title)
    for block in blocks:
        kind = block['kind']
        if kind == 'table':
            report.table([[plain(cell) for cell in row] for row in block['rows']], caption=plain(block['caption']))
            table = report.document.tables[-1]
            for i, row in enumerate(block['rows']):
                for j, cell in enumerate(row):
                    paragraph = table.cell(i, j).paragraphs[0]
                    paragraph.clear()
                    word_runs(paragraph, cell)
                    if i == 0:
                        for run in paragraph.runs:
                            run.bold = True
            continue
        style = f"Heading {block['level']}" if kind == 'heading' else 'Report Reference' if kind == 'reference' else 'Normal'
        paragraph = report.document.add_paragraph(style=style)
        word_runs(paragraph, block['parts'])
        if kind == 'quote':
            from docx.shared import Mm
            paragraph.paragraph_format.left_indent = Mm(7)
            paragraph.paragraph_format.right_indent = Mm(7)
    return report.build()


def pdf_markup(parts):
    result = ''
    for part in parts:
        content = escape(part['text'])
        if part['italic']:
            content = '<i>' + content + '</i>'
        if part['bold']:
            content = '<b>' + content + '</b>'
        if part['url']:
            href = escape(part['url'], {'"': '&quot;'})
            content = f'<link href="{href}">{content}</link>'
        result += content
    return result


def build_pdf(text, output, template, language='en', regular_font=None, bold_font=None, italic_font=None, bold_italic_font=None, regular_subfont_index=0, bold_subfont_index=0):
    from reportlab.platypus import LongTable, TableStyle, Spacer
    from reportlab.lib import colors
    title, blocks = parse(text)
    if not regular_font or not bold_font:
        raise ValueError('PDF requires explicit regular and bold font files')
    has_italic = any(p['italic'] for b in blocks for p in b.get('parts', []) + [p for row in b.get('rows', []) for cell in row for p in cell])
    if has_italic and not italic_font:
        raise ValueError('The manuscript uses italics; supply --italic-font to preserve them')
    module = load_template(template)
    Paragraph = module.Paragraph
    report = module.ScientificReport(output, plain(title), regular_font, bold_font, italic_font=italic_font, bold_italic_font=bold_italic_font, language=language, profile='academic', regular_subfont_index=regular_subfont_index, bold_subfont_index=bold_subfont_index)
    report.add_front_matter()
    for block in blocks:
        kind = block['kind']
        if kind == 'table':
            report.story.append(Paragraph(pdf_markup(block['caption']), report.styles['caption'].clone('TableCaption', keepWithNext=True)))
            head_style = report.styles['table_head'].clone('AcademicTableHead', textColor=colors.black)
            rows = [[Paragraph(pdf_markup(cell), head_style if index == 0 else report.styles['table']) for cell in row] for index, row in enumerate(block['rows'])]
            table = LongTable(rows, colWidths=[report.doc.width / len(rows[0])] * len(rows[0]), repeatRows=1)
            table.setStyle(TableStyle([('FONTNAME', (0, 0), (-1, -1), 'ResearchRegular'), ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#eeeeee')), ('LINEBELOW', (0, 0), (-1, 0), .5, colors.black), ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]))
            report.story.extend([table, Spacer(1, 8)])
            continue
        key = ('h1' if block['level'] == 1 else 'h2') if kind == 'heading' else 'reference' if kind == 'reference' else 'body'
        style = report.styles[key]
        if kind == 'quote':
            style = style.clone('BlockQuote', leftIndent=20, rightIndent=20)
        report.story.append(Paragraph(pdf_markup(block['parts']), style))
    return report.build()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('manuscript', type=Path)
    parser.add_argument('output', type=Path)
    parser.add_argument('--language', choices=['en', 'zh', 'mixed'], default='en')
    parser.add_argument('--require-references', action='store_true')
    parser.add_argument('--latin-font', default='Times New Roman', help='Word font family')
    parser.add_argument('--cjk-font', default='Noto Serif SC', help='Word CJK font family')
    parser.add_argument('--regular-subfont-index', type=int, default=0)
    parser.add_argument('--bold-subfont-index', type=int, default=0)
    for name in ['regular', 'bold', 'italic', 'bold-italic']:
        parser.add_argument(f'--{name}-font', type=Path)
    args = parser.parse_args()
    text = args.manuscript.read_text(encoding='utf-8')
    _, blocks = parse(text)
    if args.require_references and not any(b['kind'] == 'reference' for b in blocks):
        parser.error('A References / 参考文献 heading with separate entries is required')
    if args.output.resolve() == args.manuscript.resolve():
        parser.error('Output must not overwrite manuscript')
    assets = Path(__file__).resolve().parents[1] / 'assets'
    if args.output.suffix.lower() == '.docx':
        build_docx(text, args.output, assets / 'python-docx-scientific-template.py', args.language, args.latin_font, args.cjk_font)
    elif args.output.suffix.lower() == '.pdf':
        build_pdf(text, args.output, assets / 'reportlab-scientific-template.py', args.language, args.regular_font, args.bold_font, args.italic_font, args.bold_italic_font, args.regular_subfont_index, args.bold_subfont_index)
    else:
        parser.error('Output must be .docx or .pdf; use the corresponding Skill package')
    print(json.dumps({'source_sha256': hashlib.sha256(args.manuscript.read_bytes()).hexdigest(), 'output_sha256': hashlib.sha256(args.output.read_bytes()).hexdigest(), 'output': str(args.output), 'visual_qa': 'not_performed'}, ensure_ascii=False))


if __name__ == '__main__':
    main()
