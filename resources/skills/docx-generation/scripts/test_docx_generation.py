#!/usr/bin/env python3
"""Behavioral checks for the bundled scientific DOCX template and QA gate."""

from __future__ import annotations

import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]


def load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


class ScientificDocxTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        try:
            from docx import Document
        except ImportError as exc:
            raise unittest.SkipTest("python-docx is unavailable") from exc
        cls.Document = staticmethod(Document)
        cls.template = load(
            "scientific_docx_template",
            ROOT / "assets" / "python-docx-scientific-template.py",
        )
        cls.gate = load("docx_quality_gate", ROOT / "scripts" / "docx_quality_gate.py")

    def test_chinese_report_preserves_heading_table_caption_and_hyperlinked_doi(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "review.docx"
            report = self.template.ScientificDocxReport(
                output,
                "黄河中游遗址证据综述",
                language="zh",
                subtitle="限定范围的叙述性综述",
            )
            report.add_front_matter()
            report.heading("方法", 1)
            report.paragraph("纳入两项已核查的原始研究；以下结论只适用于所述遗址。")
            report.table(
                [["遗址", "证据"], ["甲", "测年与植物遗存"]],
                caption="表 1 遗址证据与范围。",
            )
            report.heading("参考文献", 1)
            report.reference("Author (2024). Verified study.", doi="10.1234/example")
            report.build()

            document = self.Document(output)
            metrics, findings = self.gate.inspect_docx(
                output, language="zh", require_doi_links=True
            )

        self.assertEqual(document.paragraphs[0].style.name, "Title")
        self.assertFalse(document.styles["Title"].element.xpath("./w:pPr/w:pBdr"))
        self.assertFalse(
            document.styles["Title"].element.xpath("./w:rPr/w:color/@w:themeColor")
        )
        self.assertEqual(
            document.sections[0]
            .footer.paragraphs[0]
            ._p.xpath(".//w:fldSimple/@w:instr"),
            ["PAGE"],
        )
        self.assertTrue(
            any(
                p.text == "方法" and p.style.name == "Heading 1"
                for p in document.paragraphs
            )
        )
        self.assertEqual(len(document.tables), 1)
        self.assertTrue(any(p.text.startswith("表 1") for p in document.paragraphs))
        self.assertGreater(metrics["cjk_characters"], 10)
        self.assertEqual(findings, [])

    def test_english_report_keeps_figures_with_captions_and_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "report.docx"
            image_path = Path(temp_dir) / "plot.png"
            try:
                from PIL import Image
            except ImportError as exc:
                raise unittest.SkipTest("Pillow is unavailable") from exc
            Image.new("RGB", (800, 480), "white").save(image_path)
            report = self.template.ScientificDocxReport(
                output, "Dataset coverage", language="en"
            )
            report.add_front_matter()
            report.heading("Results", 1)
            report.paragraph("The observed count was 18 of 100 records (18%).")
            report.figure(
                image_path, "Figure 1. Coverage among 100 records; bars show counts."
            )
            report.build()

            document = self.Document(output)
            metrics, findings = self.gate.inspect_docx(output, language="en")

        from docx.enum.text import WD_LINE_SPACING

        self.assertEqual(document.core_properties.title, "Dataset coverage")
        self.assertEqual(document.core_properties.author, "Mobius Science")
        self.assertEqual(len(document.inline_shapes), 1)
        picture_paragraph = next(
            p for p in document.paragraphs if p._p.xpath(".//w:drawing")
        )
        self.assertEqual(
            picture_paragraph.paragraph_format.line_spacing_rule, WD_LINE_SPACING.SINGLE
        )
        self.assertTrue(
            any(p.text.startswith("Figure 1.") for p in document.paragraphs)
        )
        self.assertGreater(metrics["english_words"], 10)
        self.assertEqual(findings, [])

    def test_quality_gate_finds_unlinked_doi_and_length_failure(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "bad.docx"
            document = self.Document()
            document.add_paragraph("Short title", style="Title")
            document.add_paragraph("https://doi.org/10.1234/unlinked")
            document.save(output)

            _, findings = self.gate.inspect_docx(
                output, language="en", minimum_words=50, require_doi_links=True
            )

        self.assertEqual(
            {item.code for item in findings}, {"short-report", "unlinked-doi"}
        )

    def test_render_does_not_accept_a_stale_pdf_from_an_earlier_revision(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            directory = Path(temp_dir)
            output = directory / "report.docx"
            self.Document().save(output)
            qa_dir = directory / "qa"
            qa_dir.mkdir()
            (qa_dir / "report.pdf").write_bytes(b"old proof")
            fake_soffice = directory / "soffice"
            fake_soffice.write_text("#!/bin/sh\nexit 0\n", encoding="utf-8")
            fake_soffice.chmod(0o755)

            with self.assertRaisesRegex(RuntimeError, "DOCX render failed"):
                self.gate.render_docx(output, qa_dir, fake_soffice)

    def test_bundled_cjk_font_config_and_missing_font_detection(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            environment = self.gate.fontconfig_environment(Path(temp_dir))
            config = Path(environment["FONTCONFIG_FILE"]).read_text(encoding="utf-8")
        self.assertIn("NotoSerifSC-Regular.otf", str(self.gate.BUNDLED_CJK_FONT))
        self.assertIn("/assets/fonts", config)
        self.assertTrue(
            self.gate.missing_expected_cjk_font("LiberationSerif\n", "Noto Serif SC")
        )
        self.assertFalse(
            self.gate.missing_expected_cjk_font(
                "BAAAAA+NotoSerifSC-Regular\n", "Noto Serif SC"
            )
        )


if __name__ == "__main__":
    unittest.main()
