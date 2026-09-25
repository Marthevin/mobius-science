"""Small, editable Word report template for scientific manuscripts.

Copy this file into the report workspace before adapting it. The document
source, figures, and citation metadata remain the rebuild authority.
Requires python-docx; render the final DOCX for page-by-page review.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


INK = RGBColor(22, 35, 46)
MUTED = RGBColor(82, 99, 111)
ACCENT = RGBColor(8, 105, 118)
BORDER = "D9D9D9"
PALE = "F2F7F8"


def _set_font(
    style, latin: str, east_asia: str, size: float, *, bold: bool = False
) -> None:
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = INK
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for family in ("ascii", "hAnsi", "cs"):
        r_fonts.set(qn(f"w:{family}"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    # python-docx's built-in styles keep theme attributes alongside these
    # explicit values. Word/LibreOffice may prefer the theme and silently use
    # an unrelated fallback font or color, including missing CJK glyphs.
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        r_fonts.attrib.pop(qn(f"w:{attribute}"), None)
    color = r_pr.find(qn("w:color"))
    if color is not None:
        for attribute in ("themeColor", "themeTint", "themeShade"):
            color.attrib.pop(qn(f"w:{attribute}"), None)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _set_cell_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), BORDER)
        borders.append(edge)


def _set_cell_padding(cell, value: int = 105) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side in ("top", "left", "bottom", "right"):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)


def _keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "086976")
    properties.append(color)
    run.append(properties)
    content = OxmlElement("w:t")
    content.text = text
    run.append(content)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


class ScientificDocxReport:
    """Compose one scientific report with real Word styles and editable content."""

    def __init__(
        self,
        output: str | Path,
        title: str,
        *,
        language: str = "en",
        subtitle: str = "",
        metadata_line: str = "",
        author: str = "Mobius Science",
        latin_font: str = "Times New Roman",
        cjk_font: str = "Noto Serif SC",
    ) -> None:
        if language not in {"en", "zh", "mixed"}:
            raise ValueError("language must be en, zh, or mixed")
        self.output = Path(output)
        self.title = title
        self.subtitle = subtitle
        self.metadata_line = metadata_line
        self.language = language
        self.document = Document()
        self.document.core_properties.title = title
        self.document.core_properties.author = author
        self.document.core_properties.language = (
            "zh-CN" if language == "zh" else "en-US"
        )
        self._configure_styles(latin_font, cjk_font)

    def _configure_styles(self, latin_font: str, cjk_font: str) -> None:
        section = self.document.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(22)
        section.bottom_margin = Mm(21)
        section.left_margin = Mm(23)
        section.right_margin = Mm(23)
        section.footer_distance = Mm(10)

        styles = self.document.styles
        body = styles["Normal"]
        _set_font(body, latin_font, cjk_font, 10.5 if self.language == "zh" else 11)
        body.paragraph_format.line_spacing = Pt(17 if self.language == "zh" else 16)
        body.paragraph_format.space_after = Pt(7)
        body.paragraph_format.widow_control = True

        title = styles["Title"]
        _set_font(
            title, latin_font, cjk_font, 20 if self.language == "zh" else 19, bold=True
        )
        title_p_pr = title.element.get_or_add_pPr()
        title_border = title_p_pr.find(qn("w:pBdr"))
        if title_border is not None:
            title_p_pr.remove(title_border)
        title.paragraph_format.line_spacing = Pt(27)
        title.paragraph_format.space_after = Pt(11)

        subtitle = styles["Subtitle"]
        _set_font(subtitle, latin_font, cjk_font, 11)
        subtitle.font.color.rgb = MUTED
        subtitle.paragraph_format.space_after = Pt(10)

        for level, size in ((1, 14), (2, 11.5), (3, 10.5)):
            style = styles[f"Heading {level}"]
            _set_font(style, latin_font, cjk_font, size, bold=True)
            style.paragraph_format.space_before = Pt(14 if level == 1 else 10)
            style.paragraph_format.space_after = Pt(5)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.keep_together = True

        for name, size, after in (
            ("Report Metadata", 9, 11),
            ("Report Caption", 9, 9),
            ("Report Reference", 9.5, 5),
        ):
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = body
            _set_font(style, latin_font, cjk_font, size)
            style.paragraph_format.space_after = Pt(after)
        styles["Report Metadata"].font.color.rgb = MUTED
        styles["Report Caption"].font.color.rgb = MUTED
        styles["Report Caption"].paragraph_format.keep_together = True
        styles["Report Reference"].paragraph_format.left_indent = Mm(5)
        styles["Report Reference"].paragraph_format.first_line_indent = Mm(-5)
        styles["Report Reference"].paragraph_format.line_spacing = Pt(14)

        footer_style = styles.add_style("Report Footer", WD_STYLE_TYPE.PARAGRAPH)
        footer_style.base_style = body
        _set_font(footer_style, latin_font, cjk_font, 9)
        footer_style.font.color.rgb = MUTED
        footer = section.footer.paragraphs[0]
        footer.style = footer_style
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        page_field = OxmlElement("w:fldSimple")
        page_field.set(qn("w:instr"), "PAGE")
        page_run = OxmlElement("w:r")
        page_text = OxmlElement("w:t")
        page_text.text = "1"
        page_run.append(page_text)
        page_field.append(page_run)
        footer._p.append(page_field)

    def add_front_matter(self) -> None:
        paragraph = self.document.add_paragraph(self.title, style="Title")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if self.subtitle:
            self.document.add_paragraph(self.subtitle, style="Subtitle")
        if self.metadata_line:
            self.document.add_paragraph(self.metadata_line, style="Report Metadata")

    def heading(self, text: str, level: int = 1) -> None:
        if level not in {1, 2, 3}:
            raise ValueError("heading level must be 1, 2, or 3")
        self.document.add_heading(text, level=level)

    def paragraph(self, text: str) -> None:
        self.document.add_paragraph(text, style="Normal")

    def table(
        self,
        rows: list[list[object]],
        *,
        caption: str,
        widths_mm: list[float] | None = None,
        repeat_header: bool = True,
    ) -> None:
        if not rows or not rows[0] or any(len(row) != len(rows[0]) for row in rows):
            raise ValueError("Table requires equally sized non-empty rows")
        if widths_mm and (len(widths_mm) != len(rows[0]) or sum(widths_mm) > 164):
            raise ValueError("Table widths must fit the 164 mm content frame")
        caption_paragraph = self.document.add_paragraph(caption, style="Report Caption")
        caption_paragraph.paragraph_format.keep_with_next = True
        table = self.document.add_table(rows=len(rows), cols=len(rows[0]))
        table.autofit = widths_mm is None
        for row_index, row in enumerate(rows):
            word_row = table.rows[row_index]
            _keep_row_together(word_row)
            if row_index == 0 and repeat_header:
                _repeat_header(word_row)
            for column_index, value in enumerate(row):
                cell = word_row.cells[column_index]
                cell.text = str(value)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if widths_mm:
                    cell.width = Mm(widths_mm[column_index])
                _set_cell_border(cell)
                _set_cell_padding(cell)
                if row_index == 0 and repeat_header:
                    _set_cell_shading(cell, "DDECEF")
                elif row_index % 2 == 0:
                    _set_cell_shading(cell, PALE)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = Pt(13)
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        if row_index == 0 and repeat_header:
                            run.bold = True
        self.document.add_paragraph().paragraph_format.space_after = Pt(2)

    def figure(self, path: str | Path, caption: str, *, width_mm: float = 150) -> None:
        image_path = Path(path)
        if not image_path.is_file():
            raise FileNotFoundError(image_path)
        if not 0 < width_mm <= 164:
            raise ValueError("Figure width must fit the 164 mm content frame")
        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.add_run().add_picture(str(image_path), width=Mm(width_mm))
        self.document.add_paragraph(caption, style="Report Caption")

    def reference(self, citation: str, *, doi: str | None = None) -> None:
        paragraph = self.document.add_paragraph(style="Report Reference")
        paragraph.add_run(citation)
        if doi:
            identifier = doi.strip().removeprefix("https://doi.org/")
            if not identifier.startswith("10.") or "/" not in identifier:
                raise ValueError(f"Invalid DOI: {doi}")
            url = f"https://doi.org/{identifier}"
            paragraph.add_run(" ")
            _add_hyperlink(paragraph, url, url)

    def build(self) -> Path:
        self.output.parent.mkdir(parents=True, exist_ok=True)
        self.document.save(self.output)
        return self.output
