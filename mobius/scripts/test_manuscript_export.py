"""Regression fixtures for actual document defects, not scholarly quality scores."""
import importlib.util
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SPEC = importlib.util.spec_from_file_location('manuscript_export', ROOT / 'resources/skills/_shared/scientific-report/build_manuscript.py')


class ManuscriptExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.builder = importlib.util.module_from_spec(SPEC)
        SPEC.loader.exec_module(cls.builder)

    def test_merged_references_are_rejected_before_writing(self):
        with self.assertRaisesRegex(ValueError, 'blank line'):
            self.builder.parse('# Title\n\n## References\n\nAuthor A. *Book A*.\nAuthor B. *Book B*.')

    def test_does_not_silently_drop_unsupported_content(self):
        for text in ['<script>bad</script>', '```python\nx = 1\n```', '![figure](https://example.invalid/image.png)']:
            with self.subTest(text=text), self.assertRaises(ValueError):
                self.builder.parse('# Title\n\n' + text)

    def test_docx_has_separate_references_true_italics_link_spacing_and_no_theme(self):
        from docx import Document
        text = '# A Scholarly Essay\n\n## An Argument\n\nA *specific* interpretation & its **rival**.\n\n## References\n\nAuthor A. (2024). *Book A*. https://doi.org/10.1234/test\n\nAuthor B. (2025). *Book B*.\n'
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'essay.docx'
            self.builder.build_docx(text, path, ROOT / 'resources/skills/docx-generation/assets/python-docx-scientific-template.py')
            doc = Document(path)
        self.assertEqual([p.text for p in doc.paragraphs if p.style.name == 'Heading 1'], ['An Argument', 'References'])
        refs = [p for p in doc.paragraphs if p.style.name == 'Report Reference']
        self.assertEqual(len(refs), 2)
        self.assertTrue(any(r.italic and r.text == 'Book A' for r in refs[0].runs))
        self.assertIn('. https://doi.org/', ''.join(refs[0]._p.xpath('.//w:t/text()')))
        self.assertFalse(doc.styles['Title'].element.xpath('./w:pPr/w:pBdr'))
        self.assertFalse(doc.sections[0]._sectPr.xpath('./w:docGrid'))

    def test_chinese_and_media_text_is_preserved(self):
        from docx import Document
        text = '# 中文研究计划\n\n## 研究问题\n\n中文媒体是本研究的分析对象，不作为人口估计。\n\n## 参考文献\n\n作者（2024）。报道标题。\n'
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'proposal.docx'
            self.builder.build_docx(text, path, ROOT / 'resources/skills/docx-generation/assets/python-docx-scientific-template.py', language='zh')
            doc = Document(path)
        self.assertIn('中文媒体', '\n'.join(p.text for p in doc.paragraphs))

    def test_bare_doi_links_keep_exact_visible_punctuation(self):
        from docx import Document
        text = '# Title\n\n## References\n\n作者（2024）。题目。DOI: 10.1234/test(1).\n'
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'refs.docx'
            self.builder.build_docx(text, path, ROOT / 'resources/skills/docx-generation/assets/python-docx-scientific-template.py', language='zh')
            doc = Document(path)
        ref = doc.paragraphs[-1]
        self.assertEqual(''.join(ref._p.xpath('.//w:t/text()')), '作者（2024）。题目。DOI: 10.1234/test(1).')
        self.assertTrue(any(r.target_ref == 'https://doi.org/10.1234/test(1)' for r in doc.part.rels.values()))

    def test_table_is_editable_and_requires_a_numbered_caption(self):
        from docx import Document
        table = '| Question | Evidence |\n|---|---|\n| Scope | *Inspected* passage |\n'
        with self.assertRaisesRegex(ValueError, 'caption'):
            self.builder.parse('# Title\n\n' + table)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'plan.docx'
            self.builder.build_docx('# Title\n\nTable 1. Design mapping.\n\n' + table, path, ROOT / 'resources/skills/docx-generation/assets/python-docx-scientific-template.py')
            doc = Document(path)
        self.assertEqual(len(doc.tables), 1)
        self.assertTrue(any(r.italic for r in doc.tables[0].cell(1, 1).paragraphs[0].runs))

    def test_pdf_preserves_italic_face_text_and_clickable_reference(self):
        from pypdf import PdfReader
        font_dir = Path('/System/Library/Fonts/Supplemental')
        regular = font_dir / 'Times New Roman.ttf'
        if not regular.exists():
            self.skipTest('macOS font fixture unavailable; run the rendered release fixtures')
        text = '# A Scholarly Essay\n\n## Argument\n\nA *specific* interpretation & its **rival**.\n\n## References\n\nAuthor A. *Book A*. https://doi.org/10.1234/test\n'
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'essay.pdf'
            self.builder.build_pdf(text, path, ROOT / 'resources/skills/pdf-report-generation/assets/reportlab-scientific-template.py', regular_font=regular, bold_font=font_dir / 'Times New Roman Bold.ttf', italic_font=font_dir / 'Times New Roman Italic.ttf', bold_italic_font=font_dir / 'Times New Roman Bold Italic.ttf')
            pdf = PdfReader(path)
            page = pdf.pages[0]
            self.assertIn('interpretation & its rival', page.extract_text())
            self.assertNotIn('*Book', page.extract_text())
            fonts = [str(f.get_object().get('/BaseFont')) for f in page['/Resources']['/Font'].values()]
            self.assertTrue(any('Italic' in f for f in fonts), fonts)
            self.assertEqual(page['/Annots'][0].get_object()['/A']['/URI'], 'https://doi.org/10.1234/test')

    def test_chinese_pdf_does_not_start_lines_with_closing_punctuation(self):
        import subprocess, sys
        import textwrap
        script = textwrap.dedent("""
            import importlib.util, tempfile
            from pathlib import Path
            ROOT = Path('__ROOT__')
            spec = importlib.util.spec_from_file_location('builder', ROOT / 'resources/skills/_shared/scientific-report/build_manuscript.py')
            builder = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(builder)
            font = Path('/System/Library/Fonts/Supplemental/Songti.ttc')
            if not font.exists():
                raise SystemExit(0)
            module = builder.load_template(ROOT / 'resources/skills/pdf-report-generation/assets/reportlab-scientific-template.py')
            with tempfile.TemporaryDirectory() as directory:
                report = module.ScientificReport(Path(directory) / 'zh.pdf', '中文', font, font,
                    language='zh', regular_subfont_index=6, bold_subfont_index=1)
                for mark in ('，', '；', '：', '！', '？', '）', '”', '’', '”，', '》。'):
                    for rich in (False, True):
                        source = ('<b>研究</b>结果' if rich else '研究结果') + mark + '可以检验。'
                        paragraph = module.Paragraph(source, report.styles['body'])
                        paragraph.wrap(42.1, 500)
                        lines = paragraph.blPara.lines
                        text = [''.join(line[1]) for line in lines] if paragraph.blPara.kind == 0 else [''.join(f.text for f in line.words) for line in lines]
                        assert all(line[0] not in '，；：！？）”’。》' for line in text), text
                        assert ''.join(text) == '研究结果' + mark + '可以检验。'
                raw = '研究结果”，可以检验。' * 60
                paragraph = module.Paragraph('<link href="https://example.org/">' + raw + '</link>', report.styles['body'])
                delivered = ''
                for index in range(100):
                    # The remaining fragment is reflowed into alternating frames.
                    width = 73.5 if index % 2 else 84
                    paragraph.wrap(width, 10000)
                    pieces = paragraph.split(width, 51)
                    assert pieces
                    delivered += ''.join(w.text for line in pieces[0].blPara.lines for w in line.words)
                    assert all(w.link for line in pieces[0].blPara.lines for w in line.words if w.text)
                    if len(pieces) == 1:
                        break
                    paragraph = pieces[1]
                assert delivered == raw, (len(delivered), len(raw))
        """)
        subprocess.run([sys.executable, '-c', script.replace('__ROOT__', str(ROOT))], check=True)


if __name__ == '__main__':
    unittest.main()
